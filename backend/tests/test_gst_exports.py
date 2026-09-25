from openpyxl import load_workbook
import io

from tests.conftest import auth_headers, login


async def _setup(client, headers, company_id, financial_year_id):
    await client.post(
        f"/api/v1/gst/profile?company_id={company_id}",
        json={"gstin": "27AAPFU0939F1ZV", "legal_name": "Test Co"},
        headers=headers,
    )
    period = await client.post(
        f"/api/v1/gst/return-periods?company_id={company_id}",
        json={"financial_year_id": str(financial_year_id), "year": 2025, "month": 4},
        headers=headers,
    )
    return period.json()["data"]


class TestGSTExports:
    async def test_gstr1_csv_export(self, client, company_a_with_admin, financial_year_a):
        company, admin = company_a_with_admin
        data = await login(client, admin.email, "TestPass1!")
        headers = auth_headers(data["access_token"])
        period = await _setup(client, headers, company.id, financial_year_a.id)

        response = await client.get(
            f"/api/v1/gst/return-periods/{period['id']}/reports/gstr1?company_id={company.id}&format=csv",
            headers=headers,
        )
        assert response.status_code == 200, response.text
        assert response.headers["content-type"].startswith("text/csv")
        body = response.content.decode("utf-8-sig")
        assert "GSTR-1 Preparation" in body
        assert "27AAPFU0939F1ZV" in body
        assert "B2B" in body

    async def test_gstr1_xlsx_export_has_expected_sheets(self, client, company_a_with_admin, financial_year_a):
        company, admin = company_a_with_admin
        data = await login(client, admin.email, "TestPass1!")
        headers = auth_headers(data["access_token"])
        period = await _setup(client, headers, company.id, financial_year_a.id)

        response = await client.get(
            f"/api/v1/gst/return-periods/{period['id']}/reports/gstr1?company_id={company.id}&format=xlsx",
            headers=headers,
        )
        assert response.status_code == 200, response.text
        workbook = load_workbook(io.BytesIO(response.content))
        assert "Cover" in workbook.sheetnames
        assert "B2B" in workbook.sheetnames
        assert "Exports" in workbook.sheetnames
        assert "HSN Summary" in workbook.sheetnames

    async def test_gstr1_export_includes_table6_exports(self, client, company_a_with_admin, financial_year_a):
        company, admin = company_a_with_admin
        data = await login(client, admin.email, "TestPass1!")
        headers = auth_headers(data["access_token"])
        period = await _setup(client, headers, company.id, financial_year_a.id)

        # Create export customer
        cust_resp = await client.post(
            f"/api/v1/accounting/customers?company_id={company.id}",
            json={"name": "Global Export Corp", "is_export": True},
            headers=headers,
        )
        customer_id = cust_resp.json()["data"]["id"]

        # Create export invoice
        inv_resp = await client.post(
            f"/api/v1/accounting/sales-invoices?company_id={company.id}",
            json={
                "financial_year_id": str(financial_year_a.id),
                "customer_id": customer_id,
                "invoice_number": "EXP-TBL6-01",
                "invoice_date": "2025-04-15",
                "place_of_supply_state_code": "96",
                "export_type": "WITH_PAYMENT",
                "shipping_bill_number": "SB-445566",
                "shipping_bill_date": "2025-04-16",
                "port_code": "INMAA1",
                "items": [{"quantity": 1, "unit_price": 75000, "igst_rate": 18}],
            },
            headers=headers,
        )
        assert inv_resp.status_code == 201
        inv_id = inv_resp.json()["data"]["id"]
        await client.post(
            f"/api/v1/accounting/sales-invoices/{inv_id}/post?company_id={company.id}",
            headers=headers,
        )

        # CSV export
        csv_resp = await client.get(
            f"/api/v1/gst/return-periods/{period['id']}/reports/gstr1?company_id={company.id}&format=csv",
            headers=headers,
        )
        assert csv_resp.status_code == 200
        csv_body = csv_resp.content.decode("utf-8-sig")
        assert "Exports" in csv_body
        assert "EXP-TBL6-01" in csv_body
        assert "SB-445566" in csv_body
        assert "WITH_PAYMENT" in csv_body

        # XLSX export
        xlsx_resp = await client.get(
            f"/api/v1/gst/return-periods/{period['id']}/reports/gstr1?company_id={company.id}&format=xlsx",
            headers=headers,
        )
        assert xlsx_resp.status_code == 200
        wb = load_workbook(io.BytesIO(xlsx_resp.content))
        assert "Exports" in wb.sheetnames

    async def test_gstr3b_export(self, client, company_a_with_admin, financial_year_a):
        company, admin = company_a_with_admin
        data = await login(client, admin.email, "TestPass1!")
        headers = auth_headers(data["access_token"])
        period = await _setup(client, headers, company.id, financial_year_a.id)

        response = await client.get(
            f"/api/v1/gst/return-periods/{period['id']}/reports/gstr3b?company_id={company.id}",
            headers=headers,
        )
        assert response.status_code == 200, response.text
        assert "GSTR-3B Preparation" in response.content.decode("utf-8-sig")

    async def test_reconciliation_export_requires_a_run(self, client, company_a_with_admin, financial_year_a):
        company, admin = company_a_with_admin
        data = await login(client, admin.email, "TestPass1!")
        headers = auth_headers(data["access_token"])
        period = await _setup(client, headers, company.id, financial_year_a.id)

        response = await client.get(
            f"/api/v1/gst/return-periods/{period['id']}/reports/reconciliation?company_id={company.id}",
            headers=headers,
        )
        assert response.status_code == 404
        assert response.json()["error"]["code"] == "GST_RECONCILIATION_NOT_FOUND"

    async def test_itc_export_requires_a_reconciliation_run(self, client, company_a_with_admin, financial_year_a):
        company, admin = company_a_with_admin
        data = await login(client, admin.email, "TestPass1!")
        headers = auth_headers(data["access_token"])
        period = await _setup(client, headers, company.id, financial_year_a.id)

        response = await client.get(
            f"/api/v1/gst/return-periods/{period['id']}/reports/itc?company_id={company.id}",
            headers=headers,
        )
        assert response.status_code == 404
        assert response.json()["error"]["code"] == "GST_RECONCILIATION_NOT_FOUND"

    async def test_itc_export_after_reconciliation(self, client, company_a_with_admin, financial_year_a):
        company, admin = company_a_with_admin
        data = await login(client, admin.email, "TestPass1!")
        headers = auth_headers(data["access_token"])
        period = await _setup(client, headers, company.id, financial_year_a.id)

        await client.post(
            f"/api/v1/gst/return-periods/{period['id']}/reconciliation?company_id={company.id}",
            headers=headers,
        )

        response = await client.get(
            f"/api/v1/gst/return-periods/{period['id']}/reports/itc?company_id={company.id}",
            headers=headers,
        )
        assert response.status_code == 200, response.text
        assert "ITC Summary" in response.content.decode("utf-8-sig")
