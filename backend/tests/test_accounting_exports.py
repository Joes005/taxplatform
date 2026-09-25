from datetime import date
from decimal import Decimal
import io

from openpyxl import load_workbook

from app.utils.export import write_csv, write_xlsx
from tests.conftest import auth_headers, create_vendor, login


class TestExportUtils:
    def test_write_csv_formatting(self):
        headers = ["Date", "Description", "Amount", "Notes"]
        rows = [
            [date(2025, 4, 1), "Opening", Decimal("1250.50"), None],
            [date(2025, 4, 2), "Receipt", Decimal("300.00"), "Cash payment"],
        ]
        csv_bytes = write_csv(headers, rows, title="Test Report")
        assert isinstance(csv_bytes, bytes)
        content = csv_bytes.decode("utf-8-sig")
        assert "Test Report" in content
        assert "2025-04-01" in content
        assert "1250.50" in content
        assert "Cash payment" in content

    def test_write_xlsx_formatting(self):
        headers = ["Date", "Description", "Amount", "Notes"]
        rows = [
            [date(2025, 4, 1), "Opening", Decimal("1250.50"), None],
            [date(2025, 4, 2), "Receipt", Decimal("300.00"), "Cash payment"],
        ]
        xlsx_bytes = write_xlsx(headers, rows, sheet_name="Ledger", title="Excel Report")
        assert isinstance(xlsx_bytes, bytes)
        wb = load_workbook(io.BytesIO(xlsx_bytes))
        assert "Ledger" in wb.sheetnames
        ws = wb["Ledger"]
        assert ws.cell(row=1, column=1).value == "Excel Report"


class TestAccountingReportsExport:
    async def test_sales_register_export_csv(
        self, client, company_a_with_admin, financial_year_a, customer_a
    ):
        company, admin = company_a_with_admin
        data = await login(client, admin.email, "TestPass1!")
        headers = auth_headers(data["access_token"])

        # Create a sales invoice
        inv_resp = await client.post(
            f"/api/v1/accounting/sales-invoices?company_id={company.id}",
            json={
                "financial_year_id": str(financial_year_a.id),
                "customer_id": str(customer_a.id),
                "invoice_number": "EXP-INV-001",
                "invoice_date": "2025-06-01",
                "items": [
                    {"quantity": 1, "unit_price": 1000, "cgst_rate": 9, "sgst_rate": 9},
                ],
            },
            headers=headers,
        )
        assert inv_resp.status_code == 201

        # Export CSV
        resp = await client.get(
            f"/api/v1/accounting/reports/sales-register/export?company_id={company.id}&format=csv",
            headers=headers,
        )
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith("text/csv")
        assert "attachment;" in resp.headers.get("content-disposition", "")
        content = resp.content.decode("utf-8-sig")
        assert "EXP-INV-001" in content
        assert "1000" in content

    async def test_sales_register_export_xlsx(
        self, client, company_a_with_admin, financial_year_a, customer_a
    ):
        company, admin = company_a_with_admin
        data = await login(client, admin.email, "TestPass1!")
        headers = auth_headers(data["access_token"])

        resp = await client.get(
            f"/api/v1/accounting/reports/sales-register/export?company_id={company.id}&format=xlsx",
            headers=headers,
        )
        assert resp.status_code == 200
        assert "spreadsheetml" in resp.headers["content-type"]
        wb = load_workbook(io.BytesIO(resp.content))
        assert "Sales Register" in wb.sheetnames

    async def test_purchase_register_export_csv_and_xlsx(
        self, client, company_a_with_admin, financial_year_a, vendor_a
    ):
        company, admin = company_a_with_admin
        data = await login(client, admin.email, "TestPass1!")
        headers = auth_headers(data["access_token"])

        # Create purchase invoice
        inv_resp = await client.post(
            f"/api/v1/accounting/purchase-invoices?company_id={company.id}",
            json={
                "financial_year_id": str(financial_year_a.id),
                "vendor_id": str(vendor_a.id),
                "invoice_number": "PINV-EXP-001",
                "invoice_date": "2025-06-15",
                "supplier_invoice_number": "SUPP-999",
                "items": [
                    {"quantity": 2, "unit_price": 500, "cgst_rate": 9, "sgst_rate": 9},
                ],
            },
            headers=headers,
        )
        assert inv_resp.status_code == 201

        # CSV
        csv_resp = await client.get(
            f"/api/v1/accounting/reports/purchase-register/export?company_id={company.id}&format=csv",
            headers=headers,
        )
        assert csv_resp.status_code == 200
        content = csv_resp.content.decode("utf-8-sig")
        assert "PINV-EXP-001" in content
        assert vendor_a.name in content
        assert "SUPP-999" in content

        # XLSX
        xlsx_resp = await client.get(
            f"/api/v1/accounting/reports/purchase-register/export?company_id={company.id}&format=xlsx",
            headers=headers,
        )
        assert xlsx_resp.status_code == 200
        wb = load_workbook(io.BytesIO(xlsx_resp.content))
        assert "Purchase Register" in wb.sheetnames

    async def test_trial_balance_export_csv_and_xlsx(
        self, client, company_a_with_admin, ledger_a
    ):
        company, admin = company_a_with_admin
        data = await login(client, admin.email, "TestPass1!")
        headers = auth_headers(data["access_token"])

        # CSV
        csv_resp = await client.get(
            f"/api/v1/accounting/reports/trial-balance/export?company_id={company.id}&format=csv",
            headers=headers,
        )
        assert csv_resp.status_code == 200
        content = csv_resp.content.decode("utf-8-sig")
        assert "Trial Balance" in content
        assert "Ledger Name" in content

        # XLSX
        xlsx_resp = await client.get(
            f"/api/v1/accounting/reports/trial-balance/export?company_id={company.id}&format=xlsx",
            headers=headers,
        )
        assert xlsx_resp.status_code == 200
        wb = load_workbook(io.BytesIO(xlsx_resp.content))
        assert "Trial Balance" in wb.sheetnames

    async def test_export_tenant_isolation(
        self, client, company_a_with_admin, company_b_with_admin
    ):
        company_a, _admin_a = company_a_with_admin
        _company_b, admin_b = company_b_with_admin
        data_b = await login(client, admin_b.email, "TestPass1!")
        headers_b = auth_headers(data_b["access_token"])

        # Company B admin cannot access Company A sales register export
        resp = await client.get(
            f"/api/v1/accounting/reports/sales-register/export?company_id={company_a.id}&format=csv",
            headers=headers_b,
        )
        assert resp.status_code == 403
