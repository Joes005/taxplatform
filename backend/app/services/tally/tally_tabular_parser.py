import csv
from datetime import date
from decimal import Decimal
import io
from typing import Any

from openpyxl import load_workbook

from app.core.exceptions import ValidationAppError
from app.services.tally.tally_models import (
    TallyFormat,
    TallyImportBatch,
    TallyLedgerRecord,
    TallyPartyRecord,
    TallyTaxBreakdown,
    TallyVoucherLine,
    TallyVoucherRecord,
    TallyVoucherType,
)
from app.services.tally.tally_xml_parser import (
    normalize_voucher_type,
    parse_tally_amount,
    parse_tally_date,
)


class TallyTabularParser:
    """Parses Tally CSV and Excel (XLSX) exports into a Canonical TallyImportBatch."""

    def __init__(self, company_id: str = "") -> None:
        self.company_id = company_id

    def parse_csv(self, content: bytes) -> TallyImportBatch:
        text = ""
        for enc in ("utf-8-sig", "utf-8", "latin-1"):
            try:
                text = content.decode(enc)
                break
            except UnicodeDecodeError:
                continue

        if not text:
            raise ValidationAppError("Could not decode CSV content", code="INVALID_FILE")

        reader = csv.DictReader(io.StringIO(text))
        if reader.fieldnames is None:
            raise ValidationAppError("CSV file has no header row", code="INVALID_FILE")

        rows = [dict(r) for r in reader]
        return self._rows_to_batch(rows, TallyFormat.CSV)

    def parse_xlsx(self, content: bytes) -> TallyImportBatch:
        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise ValidationAppError(f"Could not read Excel file: {exc}", code="INVALID_FILE") from exc

        sheet = wb.active
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
            header = [str(h).strip() if h is not None else "" for h in header_row]
        except StopIteration:
            raise ValidationAppError("Excel file has no header row", code="INVALID_FILE") from None

        rows = []
        for r in rows_iter:
            if all(cell is None for cell in r):
                continue
            rows.append({header[i]: str(r[i]).strip() if (i < len(r) and r[i] is not None) else "" for i in range(len(header))})

        return self._rows_to_batch(rows, TallyFormat.XLSX)

    def _rows_to_batch(self, rows: list[dict[str, str]], fmt: TallyFormat) -> TallyImportBatch:
        batch = TallyImportBatch(company_id=self.company_id, detected_format=fmt)
        if not rows:
            return batch

        # Standardize keys by lowercasing and stripping spaces/underscores
        normalized_rows: list[dict[str, str]] = []
        for row in rows:
            norm_r = {}
            for k, v in row.items():
                clean_k = k.lower().replace(" ", "").replace("_", "").replace("-", "")
                norm_r[clean_k] = v.strip() if isinstance(v, str) else str(v)
            normalized_rows.append(norm_r)

        # Detect format style:
        # Style A: Daybook / Voucher entries (Date, Particulars, VoucherType, VoucherNo, Debit, Credit)
        # Style B: Invoice Register (InvoiceNo, InvoiceDate, Party, Taxable, CGST, SGST, IGST, Total)

        sample = normalized_rows[0]
        has_vtype = any("vouchertype" in k or "type" in k for k in sample.keys())
        has_debit_credit = any("debit" in k for k in sample.keys()) or any("credit" in k for k in sample.keys())

        if has_vtype and has_debit_credit:
            self._parse_daybook_style(normalized_rows, batch)
        else:
            self._parse_register_style(normalized_rows, batch)

        batch.raw_record_count = len(rows)
        return batch

    def _parse_daybook_style(self, rows: list[dict[str, str]], batch: TallyImportBatch) -> None:
        # Group rows by voucher number
        grouped_vouchers: dict[str, list[dict[str, str]]] = {}
        row_order: list[str] = []

        for idx, row in enumerate(rows, start=1):
            v_no = self._val(row, ["voucherno", "vouchernumber", "invoiceno", "docno", "refno"])
            if not v_no:
                v_no = f"VCH-{idx}"
            if v_no not in grouped_vouchers:
                grouped_vouchers[v_no] = []
                row_order.append(v_no)
            grouped_vouchers[v_no].append(row)

        for v_no in row_order:
            v_rows = grouped_vouchers[v_no]
            first = v_rows[0]

            v_type_str = self._val(first, ["vouchertype", "type"]) or "Journal"
            norm_type = normalize_voucher_type(v_type_str)
            v_date = parse_tally_date(self._val(first, ["date", "voucherdate", "invoicedate"])) or date.today()
            narration = self._val(first, ["narration", "remarks", "description"])

            lines: list[TallyVoucherLine] = []
            party_name = None

            for r in v_rows:
                particulars = self._val(r, ["particulars", "ledger", "ledgername", "account", "party"])
                if not particulars:
                    continue

                dr_amt = parse_tally_amount(self._val(r, ["debit", "debitamount", "dr"]))
                cr_amt = parse_tally_amount(self._val(r, ["credit", "creditamount", "cr"]))

                is_debit = dr_amt > Decimal(0) or (dr_amt == Decimal(0) and cr_amt == Decimal(0))
                amt = dr_amt if dr_amt > Decimal(0) else cr_amt

                lines.append(
                    TallyVoucherLine(
                        ledger_name=particulars,
                        amount=amt,
                        is_debit=is_debit,
                        narration=self._val(r, ["narration", "lineitemnarration"]),
                    )
                )

                # Collect party / ledger metadata
                if norm_type in (TallyVoucherType.SALES, TallyVoucherType.PAYMENT) and is_debit and not party_name:
                    party_name = particulars
                elif norm_type in (TallyVoucherType.PURCHASE, TallyVoucherType.RECEIPT) and not is_debit and not party_name:
                    party_name = particulars

            debit_sum = sum((l.amount for l in lines if l.is_debit), Decimal(0))
            credit_sum = sum((l.amount for l in lines if not l.is_debit), Decimal(0))
            total_amt = max(debit_sum, credit_sum)

            voucher = TallyVoucherRecord(
                voucher_type=v_type_str,
                normalized_type=norm_type,
                voucher_number=v_no,
                voucher_date=v_date,
                party_name=party_name or (lines[0].ledger_name if lines else None),
                narration=narration if narration else None,
                lines=lines,
                total_amount=total_amt,
            )
            batch.vouchers.append(voucher)

    def _parse_register_style(self, rows: list[dict[str, str]], batch: TallyImportBatch) -> None:
        for idx, row in enumerate(rows, start=1):
            v_no = self._val(row, ["invoiceno", "voucherno", "vouchernumber", "billno", "docno"]) or f"INV-{idx}"
            v_date = parse_tally_date(self._val(row, ["invoicedate", "date", "voucherdate"])) or date.today()
            party = self._val(row, ["customername", "partyname", "party", "particulars", "suppliername", "vendorname"]) or "Unknown Party"
            gstin = self._val(row, ["gstin", "partygstin", "suppliergstin", "customergstin"])

            # Amounts
            taxable = parse_tally_amount(self._val(row, ["taxableamount", "taxablevalue", "subtotal", "taxable"]))
            cgst = parse_tally_amount(self._val(row, ["cgst", "cgstamount", "centraltax"]))
            sgst = parse_tally_amount(self._val(row, ["sgst", "sgstamount", "statetax"]))
            igst = parse_tally_amount(self._val(row, ["igst", "igstamount", "integratedtax"]))
            cess = parse_tally_amount(self._val(row, ["cess", "cessamount"]))
            total = parse_tally_amount(self._val(row, ["total", "totalamount", "grandtotal", "netamount", "amount"]))

            if total == Decimal(0):
                total = taxable + cgst + sgst + igst + cess

            v_type_str = self._val(row, ["vouchertype", "type"]) or "Sales"
            norm_type = normalize_voucher_type(v_type_str)

            # Build lines
            lines: list[TallyVoucherLine] = []
            if norm_type == TallyVoucherType.SALES:
                # Party line (Debit)
                lines.append(TallyVoucherLine(ledger_name=party, amount=total, is_debit=True))
                # Revenue line (Credit)
                lines.append(TallyVoucherLine(ledger_name="Sales Revenue", amount=taxable, is_debit=False))
                if cgst > 0:
                    lines.append(TallyVoucherLine(ledger_name="Output CGST", amount=cgst, is_debit=False))
                if sgst > 0:
                    lines.append(TallyVoucherLine(ledger_name="Output SGST", amount=sgst, is_debit=False))
                if igst > 0:
                    lines.append(TallyVoucherLine(ledger_name="Output IGST", amount=igst, is_debit=False))
                if cess > 0:
                    lines.append(TallyVoucherLine(ledger_name="Output CESS", amount=cess, is_debit=False))
            elif norm_type == TallyVoucherType.PURCHASE:
                # Expense line (Debit)
                lines.append(TallyVoucherLine(ledger_name="Purchase Expense", amount=taxable, is_debit=True))
                if cgst > 0:
                    lines.append(TallyVoucherLine(ledger_name="Input CGST", amount=cgst, is_debit=True))
                if sgst > 0:
                    lines.append(TallyVoucherLine(ledger_name="Input SGST", amount=sgst, is_debit=True))
                if igst > 0:
                    lines.append(TallyVoucherLine(ledger_name="Input IGST", amount=igst, is_debit=True))
                if cess > 0:
                    lines.append(TallyVoucherLine(ledger_name="Input CESS", amount=cess, is_debit=True))
                # Vendor line (Credit)
                lines.append(TallyVoucherLine(ledger_name=party, amount=total, is_debit=False))
            else:
                lines.append(TallyVoucherLine(ledger_name=party, amount=total, is_debit=True))
                lines.append(TallyVoucherLine(ledger_name="General Ledger", amount=total, is_debit=False))

            voucher = TallyVoucherRecord(
                voucher_type=v_type_str,
                normalized_type=norm_type,
                voucher_number=v_no,
                voucher_date=v_date,
                party_name=party,
                narration=self._val(row, ["narration", "remarks"]),
                lines=lines,
                total_amount=total,
                tax_breakdown=TallyTaxBreakdown(
                    taxable_amount=taxable,
                    cgst_amount=cgst,
                    sgst_amount=sgst,
                    igst_amount=igst,
                    cess_amount=cess,
                    total_tax=cgst + sgst + igst + cess,
                ),
            )
            batch.vouchers.append(voucher)

            # Register Party if customer/vendor
            if party and party != "Unknown Party":
                party_type = "CUSTOMER" if norm_type == TallyVoucherType.SALES else "VENDOR"
                batch.parties.append(
                    TallyPartyRecord(
                        name=party,
                        party_type=party_type,
                        gstin=gstin if gstin else None,
                    )
                )

    def _val(self, row: dict[str, str], candidates: list[str]) -> str:
        for c in candidates:
            if c in row and row[c]:
                return row[c]
        return ""
